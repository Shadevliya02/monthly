"""Build data.js for the dashboard from the xlsx exports.

Usage:  python3 tools/build_data.py
Needs:  pip install openpyxl

Sources (in data/):
  Order_Outbound_Inbound_2026.xlsx  base list of all orders (sheet "Orders")
  updates/*.xlsx                    later exports (sheet "Orders"), applied in file-name order,
                                    so prefix them with the date (YYYY-MM-DD_...). A row whose
                                    invoice number already exists updates that order, otherwise it
                                    is added. Orders missing from an update file are kept.
  return_dates.json                 return/resubmit cycles keyed by invoice number
"""
import datetime
import json
import pathlib
import re
from zoneinfo import ZoneInfo

import openpyxl

ROOT = pathlib.Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
BASE = DATA / "Order_Outbound_Inbound_2026.xlsx"
OUT = ROOT / "data.js"

# partners left out of the dashboard entirely (all their invoices are dropped)
EXCLUDE_PARTNERS = {"PT Sarana Komunikasi Data"}
# single invoices left out of the dashboard: invoice number -> submit date (YYYY-MM-DD) of the record
# to drop, or None to drop every record with that number
EXCLUDE_INVOICES = {
    "269/VMP-INV/FA/VI/2025": "2026-02-12",   # the returned record; the paid one (submitted 2025-11-06) stays
    "005/INV-GMT/II/2026": None,
    "031/INV-GMT/IX/2026": None,
    "JNV/26-08-1022/AJR": None,   # typo duplicate of INV/26-08-1022/AJR (submitted 2026-09-11), which stays
    "259VMP-INVFAVIII2026": None,  # typo duplicate of 259/VMP-INV/FA/VIII/2026 (paid 2026-10-13)
    "INV/26-09-1102/AR": None,     # typo duplicate of INV/26-09-1102/AJR (paid 2026-10-12)
    "INV-250901": None,            # ICG Singapore, value 10,000 in the export (USD?): held back until confirmed
}
# invoices whose value in a later export is not taken over (kept as is until confirmed)
KEEP_VALUE = {
    "INV-250908", "INV-250902", "INV-250904",   # ICG Singapore: new export shows 122,500 / 33,000 / 90,000 (USD?)
}
# partners for which only the listed invoice numbers are shown (all their other invoices are left out for now)
KEEP_ONLY = {
    "CV Divpa Solusindo": {"003/DIVPA/V/2026"},
}
# partner names to correct (name in an export -> name shown)
PARTNER_NAMES = {
    "PT Divpa Solusindo": "CV Divpa Solusindo",
}

# accepted header names per field (exports are not always consistent)
COLS = {
    "no": ["No.", "No"],
    "work": ["Type of Work"],
    "project": ["Type Project"],
    "partner": ["Partner"],
    "contract": ["Contract Number"],
    "desc": ["Description"],
    "inv": ["Invoice Number"],
    "value": ["BAST Value"],
    "status": ["Current Status", "Status Payment", "Status"],
    "submit": ["Submit Invoice Date", "Created At", "Submit Date"],
    "paid": ["Paid Date"],
    "ret": ["Return Date"],          # e.g. "Return 1 : 14 September 2026 Return 2 : 18 September 2026"
    "resub": ["Resubmit Date"],      # e.g. "Resubmit 1 : 15 September 2026 Resubmit 2 : 15 Oktober 2026"
    "reason": ["Reason"],
}
# fields an export may leave out; an update file without them keeps the existing values
OPTIONAL = {"no", "contract", "ret", "resub", "reason"}
# Indonesian month names in the inbound exports
MONTHS_ID = {"januari": "January", "februari": "February", "maret": "March", "mei": "May", "juni": "June",
             "juli": "July", "agustus": "August", "oktober": "October", "desember": "December"}


def clean(v):
    if v is None:
        return ""
    s = str(v).replace("​", "").replace("\xa0", " ")
    return re.sub(r"\s+", " ", s).strip()


def key(inv):
    return re.sub(r"\s+", "", clean(inv)).upper()


def date(v):
    if v is None or v == "":
        return None
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    s = clean(v)
    s = re.sub(r"[A-Za-z]+", lambda m: MONTHS_ID.get(m.group(0).lower(), m.group(0)), s)
    for fmt in ("%d %B %Y", "%d %b %Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            pass
    raise ValueError(f"Unrecognised date: {v!r}")


def dates_in(text):
    """All dates written like '14 September 2026' / '08 Oktober 2026' inside a free-text cell, in order."""
    return [date(m) for m in re.findall(r"\d{1,2} [A-Za-z]+ \d{4}", clean(text))]


def read(path):
    ws = openpyxl.load_workbook(path, data_only=True)["Orders"]
    header = [clean(c) for c in next(ws.iter_rows(max_row=1, values_only=True))]
    idx = {}
    for field, names in COLS.items():
        found = [header.index(n) for n in names if n in header]
        if not found:
            if field in OPTIONAL:
                continue
            raise KeyError(f"{path.name}: column for {field!r} not found (looked for {names})")
        idx[field] = found[0]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not any(c is not None for c in r):
            continue
        row = {
            "no": r[idx["no"]] if "no" in idx else None,
            "work": clean(r[idx["work"]]),
            "project": clean(r[idx["project"]]),
            "partner": clean(r[idx["partner"]]),
            "desc": clean(r[idx["desc"]]),
            "inv": clean(r[idx["inv"]]),
            "value": float(r[idx["value"]] or 0),
            "status": clean(r[idx["status"]]).lower(),
            "submit": date(r[idx["submit"]]),
            "paid": date(r[idx["paid"]]),
        }
        if "contract" in idx:
            row["contract"] = clean(r[idx["contract"]])
        if "ret" in idx:
            # return / resubmit cycles from the free-text columns, paired in order
            rets = dates_in(r[idx["ret"]])
            subs = dates_in(r[idx["resub"]]) if "resub" in idx else []
            row["cycles"] = [[d, subs[i] if i < len(subs) else None] for i, d in enumerate(rets)]
            if "reason" in idx and clean(r[idx["reason"]]):
                row["reason"] = clean(r[idx["reason"]])
        if row["status"] == "open":
            # "open" = not paid yet: returned if the last step is a return with no resubmit after it, else in review
            cyc = row.get("cycles") or []
            row["status"] = "returned" if cyc and not cyc[-1][1] else "review"
        rows.append(row)
    return rows


def main():
    orders = read(BASE)
    by_key = {}
    for o in orders:
        by_key.setdefault(key(o["inv"]), []).append(o)
    next_no = max(o["no"] for o in orders) + 1

    sources = [BASE.name]
    stats = []
    for path in sorted((DATA / "updates").glob("*.xlsx")):
        updated = added = 0
        seen = set()
        for u in read(path):
            u["status"] = u["status"] or "paid"
            if key(u["inv"]) in KEEP_VALUE:
                u.pop("value")
            matches = by_key.get(key(u["inv"]))
            if matches and len(matches) > 1:
                # several records share this number: match on the description
                same = [o for o in matches if o["desc"].lower() == u["desc"].lower()]
                matches = same or matches
            if matches and key(u["inv"]) in seen and not any(o["desc"].lower() == u["desc"].lower() for o in matches):
                matches = None   # the same number twice in one file with different descriptions: a separate invoice
            seen.add(key(u["inv"]))
            if matches:
                for o in matches:
                    o.update({k: v for k, v in u.items() if k != "no"})
                updated += 1
            else:
                u.setdefault("contract", "")
                u["no"] = next_no
                next_no += 1
                orders.append(u)
                by_key.setdefault(key(u["inv"]), []).append(u)
                added += 1
        sources.append(path.name)
        stats.append(f"{path.name}: {updated} updated, {added} added")

    names = {clean(a).lower(): b for a, b in PARTNER_NAMES.items()}
    for o in orders:
        o["partner"] = names.get(o["partner"].lower(), o["partner"])
        # older exports put inbound as Type of Work "Project Delta" (Type Project "-"); newer ones as "Inbound" / "Project Delta"
        if o["work"] == "Project Delta":
            o["work"], o["project"] = "Inbound", "Project Delta"

    excluded = {clean(p).lower() for p in EXCLUDE_PARTNERS}
    dropped = [o for o in orders if o["partner"].lower() in excluded]
    orders = [o for o in orders if o["partner"].lower() not in excluded]
    if dropped:
        stats.append(f"left out {len(dropped)} invoices from: {', '.join(sorted(EXCLUDE_PARTNERS))}")

    skip = {key(k): v for k, v in EXCLUDE_INVOICES.items()}
    def skipped(o):
        k = key(o["inv"])
        return k in skip and (skip[k] is None or skip[k] == o["submit"])
    dropped = [o for o in orders if skipped(o)]
    orders = [o for o in orders if not skipped(o)]
    if dropped:
        stats.append(f"left out {len(dropped)} single invoices: {', '.join(o['inv'] for o in dropped)}")

    keep = {clean(p).lower(): {key(i) for i in invs} for p, invs in KEEP_ONLY.items()}
    def kept(o):
        allowed = keep.get(o["partner"].lower())
        return allowed is None or key(o["inv"]) in allowed
    dropped = [o for o in orders if not kept(o)]
    orders = [o for o in orders if kept(o)]
    if dropped:
        stats.append(f"left out {len(dropped)} other invoices from: {', '.join(sorted(KEEP_ONLY))}")

    returns = {key(k): v for k, v in json.loads((DATA / "return_dates.json").read_text()).items() if not k.startswith("_")}
    for o in orders:
        # return_dates.json (recorded by hand) wins; otherwise the cycles from the latest export that has them
        o["cycles"] = returns.get(key(o["inv"]), o.get("cycles") or [])

    payload = {"source": " + ".join(sources), "built": datetime.datetime.now(ZoneInfo("Asia/Jakarta")).isoformat(timespec="minutes"), "orders": orders}
    OUT.write_text("window.PAYMENT_DATA = " + json.dumps(payload, ensure_ascii=False, separators=(",", ":")) + ";\n")
    print(f"Wrote {OUT.name}: {len(orders)} orders, {sum(o['status'] == 'paid' for o in orders)} paid, "
          f"{sum(bool(o['cycles']) for o in orders)} with return dates")
    for s in stats:
        print("  " + s)


if __name__ == "__main__":
    main()
